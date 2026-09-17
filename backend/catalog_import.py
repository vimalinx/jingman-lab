"""Offline, new-database-only import. No credentials or certificate intake.

python -m backend.catalog_import SOURCE.xlsx --output-dir PRIVATE_NEW_DIRECTORY
Without --apply, only a private preflight report is produced.
"""
import argparse
from collections import Counter
from decimal import Decimal
import hashlib
import json
import os
from pathlib import Path
import re
import time
from urllib.parse import urlsplit

from .db import Database
from .fulfillment import DEFAULT_RULES, pickup_restriction

IMAGE_HOST = 'fbbc-oss-public.oss-cn-shanghai.aliyuncs.com'

def inspect(source):
    from openpyxl import load_workbook
    wb = load_workbook(source, read_only=True, data_only=False)
    try:
        ws = wb['导出结果']
        iterator = ws.iter_rows()
        headers = [c.value for c in next(iterator)]
        required = ['商品名称*','SKU编码*','条形码*','是否称重*','零售价*','计量单位*','一级分类名','上下架','可售库存','图片地址']
        if any(headers.count(k) != 1 for k in required):
            raise ValueError('缺少必需列或存在重复表头')
        records = []
        for number, cells in enumerate(iterator, 2):
            if not any(c.value is not None for c in cells): continue
            if any(c.data_type == 'f' for c in cells):
                raise ValueError(f'第{number}行包含公式，请使用纯值导出表')
            r = dict(zip(headers, [c.value for c in cells]))
            text = lambda k: str(r.get(k) if r.get(k) is not None else '').strip()
            sku, barcode, name = text('SKU编码*'), text('条形码*'), text('商品名称*')
            if not re.fullmatch(r'\d{1,30}', sku) or not re.fullmatch(r'\d{8,14}',barcode) or not name:
                raise ValueError(f'第{number}行商品名称、SKU或条码无效')
            price, stock = Decimal(text('零售价*')), Decimal(text('可售库存'))
            if not price.is_finite() or not stock.is_finite() or not 0 < price <= 100000 or price*100 != (price*100).to_integral_value():
                raise ValueError(f'第{number}行金额或库存无效')
            if text('是否称重*') not in ('是','否') or text('上下架') not in ('上架','下架'):
                raise ValueError(f'第{number}行商品类型或上架状态无效')
            weighed = text('是否称重*') == '是'
            categories = [text(k) for k in ('一级分类名','二级分类名','三级分类名') if text(k)]
            issues = []
            if weighed: issues.append('称重结算待实现')
            if stock < 0: issues.append('源库存为负，待盘点')
            if not weighed and stock != stock.to_integral_value(): issues.append('计件商品库存非整数')
            if stock > 1000000: issues.append('库存超出预览范围')
            # Candidate review, not a legal classification based on product names.
            if any(t in name for t in ('香烟','卷烟','电子烟','雪茄')) or any('烟草' in c for c in categories):
                issues.append('线上销售范围待审核')
            if any(c in ('酒水','酒','保健品') for c in categories): issues.append('商品类目待审核')
            image = text('图片地址')
            u = urlsplit(image)
            if image and (u.scheme != 'https' or u.hostname != IMAGE_HOST or u.username or u.password or u.port not in (None,443)):
                issues.append('图片地址不在允许来源'); image = ''
            records.append(dict(row=number,source_sku=sku,barcode=barcode,name=name,
                unit=text('计量单位*'),price_cents=int(price*100),stock=str(stock),weighed=weighed,
                categories=categories,image_url=image,issues=issues,
                active=text('上下架')=='上架' and not issues,
                quantity=int(stock) if not issues else 0))
        for key in ('source_sku','barcode'):
            if len({r[key] for r in records}) != len(records): raise ValueError(f'{key} 重复，拒绝导入')
        return records
    finally: wb.close()

def summary(records):
    return dict(total=len(records),active=sum(r['active'] for r in records),
        held=sum(not r['active'] for r in records),
        missing_images=sum(not r['image_url'] for r in records),
        issues=dict(Counter(i for r in records for i in r['issues'])))

def build(source, output, apply=False):
    source, output = Path(source), Path(output)
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    records = inspect(source)
    if hashlib.sha256(source.read_bytes()).hexdigest() != before: raise ValueError('读取期间源文件变化')
    report = dict(source_sha256=before, **summary(records),review=[r for r in records if r['issues'] or not r['image_url']])
    # Refuse any existing output; never overwrite a running database or report.
    output.mkdir(mode=0o700,parents=True,exist_ok=False)
    with (output/'import-review.json').open('x',encoding='utf-8') as f:
        os.chmod(f.name,0o600);json.dump(report,f,ensure_ascii=False,indent=2)
    if not apply:return report
    db=Database(output/'jingman.sqlite3');db.initialize(seed=False)
    with db.tx() as c:
        c.execute('INSERT INTO stores(id,name,latitude,longitude,minimum_cents,delivery_cents,free_shipping_cents) VALUES(1,?,0,0,2000,200,3000)',('京漫便民超市',))
        profile=dict(address='门店地址待商家确认',
            phone='',hours='06:00–23:00',delivery_area='配送区域待商家确认',
            delivery_confirmed=False,local_preview=True,delivery_rules=DEFAULT_RULES,
            notice='商品来自门店导出快照；仅用于本地演练，不实时同步收银库存。配送门槛暂按20元起送、30元免基础配送费（超重费另计），待商家确认。')
        c.execute('INSERT INTO store_profile VALUES(1,?)',(json.dumps(profile,ensure_ascii=False),))
        c.executemany('INSERT INTO users VALUES(?,?,?,?)',[(1,'本地顾客','customer',None),(2,'测试顾客','customer',None),(10,'门店管理员','manager',1),(11,'店员','picker',1),(12,'商家自配送','rider',1)])
        c.executemany('INSERT INTO switches VALUES(?,0)',[('printer_offline',),('refund_failure',),('delivery_unavailable',)])
        c.execute('INSERT INTO import_batches VALUES(?,?,?)',(before,time.time(),json.dumps(summary(records),ensure_ascii=False)))
        categories={}
        for r in records:
            category=' / '.join(r['categories'])
            if category not in categories:
                categories[category]=len(categories)+1
                c.execute('INSERT INTO categories VALUES(?,?,?)',(categories[category],category,categories[category]))
            sku=int(r['source_sku'])
            c.execute('INSERT INTO products VALUES(?,?,?,?,?,?,?,?,?)',(sku,'sku-'+r['source_sku'],r['name'],categories[category],r['unit'],
                '门店商品快照。'+('；'.join(r['issues']) if r['issues'] else '价格和库存以当前本地资料为准。'),
                'placeholder','待审核' if r['issues'] else '门店商品',r['barcode']))
            c.execute('INSERT INTO inventory VALUES(1,?,?,?,?,0,?,1)',(sku,r['price_cents'],r['price_cents'],r['quantity'],int(r['active'])))
            reason=pickup_restriction(r['name'],r['categories'],r['weighed'],r['price_cents'])
            c.execute('INSERT INTO product_fulfillment VALUES(1,?,NULL,?,?)',(sku,int(bool(reason)),reason))
            c.execute('INSERT INTO stock_ledger(store_id,sku_id,delta_hand,delta_reserved,reason,reference,created) VALUES(1,?,?,0,?,?,?)',(sku,r['quantity'],'opening',before,time.time()))
            c.execute('INSERT INTO product_sources VALUES(?,?,?,?,?,?,?,?)',(sku,before,r['source_sku'],r['row'],r['stock'],int(r['weighed']),json.dumps(r['issues'],ensure_ascii=False),r['image_url']))
    return report

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('source',type=Path);p.add_argument('--output-dir',type=Path,required=True);p.add_argument('--apply',action='store_true')
    a=p.parse_args();os.umask(0o077)
    result=build(a.source,a.output_dir,a.apply)
    print(json.dumps({k:v for k,v in result.items() if k!='review'},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
